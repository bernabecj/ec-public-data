"""Parse Banco Central del Ecuador quarterly national accounts workbooks.

The ``tou_<bulletin>_<YYYYQQ>.xlsx`` workbook ("Oferta y Utilización de Bienes y Servicios")
has one sheet per measure. Every sheet has the same shape::

    row 9   Variables | PIB | Importaciones ... | Oferta final | ...   <- header
    row 10  Trimestres
    row 12  2000.I    | 3431205.89 | 1060230.10 | ...                  <- first quarter
    ...
    row 117 2026.II   | ...                                           <- last quarter
    row 119 Nota: (*) Incluye ... ISFLSH                              <- footnote

Parsers here are pure functions from bytes to a DataFrame, with no network or cache access.
"""

from __future__ import annotations

import io
from dataclasses import dataclass

import openpyxl
import pandas as pd

from ecuadata._parsing import parse_quarter, snake_case, to_number
from ecuadata.exceptions import ParseError


@dataclass(frozen=True)
class Measure:
    """One sheet of the workbook."""

    sheet: str
    units: str
    description: str


# (measure, seasonally_adjusted) -> sheet. Growth rates, price indices and contributions are
# only published seasonally adjusted.
MEASURES: dict[tuple[str, bool], Measure] = {
    ("current_prices", False): Measure(
        "Dem_Corr_bru", "miles de USD, precios corrientes", "Miles de USD, datos brutos"
    ),
    ("current_prices", True): Measure(
        "Dem_Corr_ajus",
        "miles de USD, precios corrientes",
        "Miles de USD, datos ajustados de estacionalidad",
    ),
    ("constant_prices", False): Measure(
        "Dem_Cad_bru",
        "miles de USD, precios encadenados de 2018",
        "Miles de USD, niveles encadenados 2018 = 100, datos brutos",
    ),
    ("constant_prices", True): Measure(
        "Dem_Cad_ajus",
        "miles de USD, precios encadenados de 2018",
        "Miles de USD, niveles encadenados 2018 = 100, datos ajustados de estacionalidad",
    ),
    ("volume_index", False): Measure(
        "Dem_Ivol_bru",
        "índice (2018 = 100)",
        "Índices de volumen encadenados 2018 = 100, datos brutos",
    ),
    ("volume_index", True): Measure(
        "Dem_Ivol_ajus",
        "índice (2018 = 100)",
        "Índices de volumen encadenados 2018 = 100, datos ajustados de estacionalidad",
    ),
    ("growth_yoy", True): Measure(
        "Dem_Ivol_ajus_vY",
        "porcentaje, variación respecto al mismo trimestre del año anterior",
        "Índices de volumen encadenados, variación interanual",
    ),
    ("growth_qoq", True): Measure(
        "Dem_Ivol_ajus_vQ",
        "porcentaje, variación respecto al trimestre anterior",
        "Índices de volumen encadenados, variación intertrimestral",
    ),
    ("price_index", True): Measure(
        "Dem_Prec_ajus",
        "índice (2018 = 100)",
        "Índices de precios 2018 = 100, datos ajustados de estacionalidad",
    ),
    ("price_index_yoy", True): Measure(
        "Dem_Prec_ajus_vY",
        "porcentaje, variación respecto al mismo trimestre del año anterior",
        "Índices de precios, variación interanual",
    ),
    ("contribution_yoy", True): Measure(
        "Dem_contr_ajus_vY",
        "puntos porcentuales del crecimiento interanual del PIB",
        "Contribuciones al crecimiento interanual",
    ),
    ("contribution_qoq", True): Measure(
        "Dem_contr_ajus_vQ",
        "puntos porcentuales del crecimiento intertrimestral del PIB",
        "Contribuciones al crecimiento intertrimestral",
    ),
}

MEASURE_NAMES = sorted({name for name, _ in MEASURES})

# Short, stable column names for the labels BCE uses. Anything not listed here is snake_cased.
COLUMNS = {
    "pib": "pib",
    "importaciones_de_bienes_y_servicios_fob": "importaciones",
    "oferta_final": "oferta_final",
    "demanda_interna": "demanda_interna",
    "gasto_de_consumo_final_hogares": "consumo_hogares",
    "gasto_de_consumo_final_gobierno_general": "consumo_gobierno",
    "fbkf": "fbkf",
    "variacion_de_existencias": "variacion_existencias",
    "exportaciones_de_bienes_y_servicios_fob": "exportaciones",
    "utilizacion_final": "utilizacion_final",
}

HEADER_LABEL = "variables"
REQUIRED_COLUMNS = ("pib",)


def measure_for(measure: str, seasonally_adjusted: bool) -> Measure:
    """Look up a sheet, with a helpful error for unknown or unpublished combinations."""
    try:
        return MEASURES[(measure, seasonally_adjusted)]
    except KeyError:
        if measure not in MEASURE_NAMES:
            raise ValueError(
                f"Medida desconocida {measure!r}. Elija una de: {', '.join(MEASURE_NAMES)}."
            ) from None
        raise ValueError(
            f"El BCE publica {measure!r} solo con ajuste estacional; "
            "llame con seasonally_adjusted=True."
        ) from None


def column_name(label: object) -> str:
    key = snake_case(str(label))
    return COLUMNS.get(key, key)


def _is_quarter(value: object) -> bool:
    try:
        parse_quarter(value)
    except (ValueError, TypeError):
        return False
    return value is not None


def parse_sheet(content: bytes, sheet: str) -> pd.DataFrame:
    """Read one sheet of the workbook into a quarterly DataFrame.

    Returns a DataFrame with a quarterly ``PeriodIndex`` named ``quarter``, float columns
    named after ``COLUMNS``, and ``attrs["original_columns"]`` mapping each column back to the
    label in the file.

    Raises:
        ParseError: the sheet is missing, or its layout is not the one described above.
    """
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises a zoo of exceptions for bad files
        raise ParseError(
            f"El archivo no es un libro XLSX legible ({type(exc).__name__}: {exc})."
        ) from exc
    try:
        if sheet not in workbook.sheetnames:
            raise ParseError(
                f"No existe la hoja {sheet!r}; el libro tiene {workbook.sheetnames}. "
                "Puede que el BCE haya renombrado sus tablas; por favor abra un issue."
            )
        rows = [list(row) for row in workbook[sheet].iter_rows(values_only=True)]
    finally:
        workbook.close()
    return _frame_from_rows(rows, sheet)


def _frame_from_rows(rows: list[list[object]], sheet: str) -> pd.DataFrame:
    header_at = next(
        (i for i, row in enumerate(rows) if row and str(row[0]).strip().lower() == HEADER_LABEL),
        None,
    )
    if header_at is None:
        raise ParseError(
            f"No se encontró la fila de encabezado {HEADER_LABEL!r} en la hoja {sheet!r}; "
            "el formato cambió."
        )
    labels = [label for label in rows[header_at][1:] if label is not None and str(label).strip()]
    if not labels:
        raise ParseError(f"El encabezado de la hoja {sheet!r} no tiene nombres de variables.")
    width = len(labels)

    quarters: list[pd.Period] = []
    values: list[list[float]] = []
    for row in rows[header_at + 1 :]:
        if not _is_quarter(row[0] if row else None):
            if quarters:
                break  # footnotes start here
            continue  # blank rows between the header and the first quarter
        quarters.append(parse_quarter(row[0]))
        cells = list(row[1 : 1 + width]) + [None] * (width - len(row[1:]))
        try:
            values.append([to_number(cell) for cell in cells])
        except (TypeError, ValueError) as exc:
            raise ParseError(
                f"Celda inesperada en la hoja {sheet!r}, fila {row[0]!r}: {exc}"
            ) from exc
    if not quarters:
        raise ParseError(f"La hoja {sheet!r} no tiene filas trimestrales bajo el encabezado.")

    index = pd.PeriodIndex(quarters, freq="Q", name="quarter")
    if not index.is_unique:
        raise ParseError(f"La hoja {sheet!r} repite el mismo trimestre dos veces.")
    if not index.is_monotonic_increasing:
        raise ParseError(f"Los trimestres de la hoja {sheet!r} no están en orden cronológico.")

    columns = [column_name(label) for label in labels]
    if len(set(columns)) != len(columns):
        raise ParseError(f"La hoja {sheet!r} tiene nombres de variables repetidos: {labels}.")
    missing = [c for c in REQUIRED_COLUMNS if c not in columns]
    if missing:
        raise ParseError(
            f"A la hoja {sheet!r} le falta(n) la(s) variable(s) {missing}; se encontró {labels}."
        )

    frame = pd.DataFrame(values, index=index, columns=columns, dtype="float64")
    frame.attrs["original_columns"] = dict(
        zip(columns, (str(x).strip() for x in labels), strict=True)
    )
    return frame


def parse_pib_trimestral(
    content: bytes, measure: str = "constant_prices", seasonally_adjusted: bool = True
) -> pd.DataFrame:
    """Parse the sheet for ``measure`` out of a ``tou`` workbook.

    ``attrs`` carries ``measure``, ``seasonally_adjusted``, ``units``, ``sheet`` and
    ``original_columns``; the caller adds provenance (``url``, ``retrieved_at`` ...).
    """
    spec = measure_for(measure, seasonally_adjusted)
    frame = parse_sheet(content, spec.sheet)
    frame.attrs.update(
        {
            "measure": measure,
            "seasonally_adjusted": seasonally_adjusted,
            "units": spec.units,
            "description": spec.description,
            "sheet": spec.sheet,
            "base_year": 2018,
        }
    )
    return frame
