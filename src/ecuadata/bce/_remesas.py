"""Workers' remittances (remesas de trabajadores) received by and sent from Ecuador.

BCE publishes the series as one workbook with a stable name::

    https://contenido.bce.fin.ec/documentos/Estadisticas/SectorExterno/BalanzaPagos/Remesas/
        Flujo_de_remesas_de_trabajadores.xlsx

Sheets follow a naming scheme: ``RR``/``RE`` = received / sent ("recibidas" / "enviadas"),
``pais`` = by country of origin or destination, ``g`` = number of transfers (giros) instead of
amounts, ``_a``/``_t`` = annual / quarterly. Every sheet has its header on row 9 and data
from row 10; total sheets have columns ``Trimestre | Monto en USD | Número de giros``, country
sheets have ``Código | Descripción | <period> ...`` with a final ``Total`` row.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from datetime import datetime

import openpyxl
import pandas as pd

from ecuadata import cache, http
from ecuadata._parsing import parse_quarter, parse_year, snake_case, to_number
from ecuadata.config import config
from ecuadata.exceptions import OfflineError, ParseError, SourceUnavailableError

PROVIDER = "bce"
SOURCE = "BCE"
URL = (
    "https://contenido.bce.fin.ec/documentos/Estadisticas/SectorExterno/BalanzaPagos/Remesas/"
    "Flujo_de_remesas_de_trabajadores.xlsx"
)
FILENAME = "Flujo_de_remesas_de_trabajadores.xlsx"
LANDING_URL = "https://contenido.bce.fin.ec/series-de-datos-remesas-de-trabajadores/"
_XLSX_MAGIC = b"PK\x03\x04"

FLOWS = {"received": "RR", "sent": "RE"}
FREQUENCIES = {"Q": ("_t", parse_quarter, "Q"), "A": ("_a", parse_year, "Y")}
TOTAL_COLUMNS = {"monto_en_usd": "monto_usd", "numero_de_giros": "numero_giros"}


@dataclass(frozen=True)
class Workbook:
    content: bytes
    url: str
    retrieved_at: datetime


def fetch(refresh: bool = False) -> Workbook:
    """Descarga el libro, o reutiliza una copia del caché más nueva que ``config.cache_ttl``."""
    if config.offline:
        cached = cache.load(PROVIDER, FILENAME)
        if cached is None:
            raise OfflineError("El libro de remesas del BCE no está en el caché.")
        return Workbook(cached.content, cached.url, cached.retrieved_at)
    if not refresh and (cached := cache.load(PROVIDER, FILENAME, max_age=config.cache_ttl)):
        return Workbook(cached.content, cached.url, cached.retrieved_at)
    response = http.get(URL)
    if response.status_code != 200 or not response.content.startswith(_XLSX_MAGIC):
        raise SourceUnavailableError(
            f"No se pudo descargar el libro de remesas del BCE desde {URL} "
            f"(HTTP {response.status_code}). Puede que el BCE lo haya movido; vea {LANDING_URL}."
        )
    stored = cache.store(PROVIDER, FILENAME, response.content, URL)
    return Workbook(stored.content, stored.url, stored.retrieved_at)


def sheet_name(flow: str, frequency: str, by: str | None, count: bool) -> str:
    if flow not in FLOWS:
        raise ValueError(f"Flujo desconocido {flow!r}; elija 'received' o 'sent'.")
    if frequency not in FREQUENCIES:
        raise ValueError(
            f"Frecuencia desconocida {frequency!r}; elija 'Q' (trimestral) o 'A' (anual)."
        )
    if by not in (None, "country"):
        raise ValueError(f"Desglose desconocido {by!r}; elija None o 'country'.")
    if count and by is None:
        raise ValueError(
            "El total ya incluye la columna 'numero_giros'; count=True solo aplica con "
            "by='country'."
        )
    name = FLOWS[flow]
    if by == "country":
        name += "pais" + ("g" if count else "")
    return name + FREQUENCIES[frequency][0]


def _rows(content: bytes, sheet: str) -> list[list[object]]:
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ParseError(
            f"El archivo no es un libro XLSX legible ({type(exc).__name__}: {exc})."
        ) from exc
    try:
        if sheet not in workbook.sheetnames:
            raise ParseError(
                f"No existe la hoja {sheet!r} en el libro de remesas "
                f"(hojas: {workbook.sheetnames[:8]}...). Puede que el BCE haya cambiado el formato."
            )
        return [list(row) for row in workbook[sheet].iter_rows(values_only=True)]
    finally:
        workbook.close()


def _header_row(rows: list[list[object]], first_label: str, sheet: str) -> int:
    for i, row in enumerate(rows):
        cells = [str(c).strip().lower() for c in row if c is not None]
        if cells and cells[0] == first_label:
            return i
    raise ParseError(f"No se encontró el encabezado {first_label!r} en la hoja {sheet!r}.")


def parse_totals(content: bytes, sheet: str, frequency: str) -> pd.DataFrame:
    """``RR_t``-style sheet -> DataFrame[period] with ``monto_usd`` and ``numero_giros``."""
    _, parse_period, freq = FREQUENCIES[frequency]
    rows = _rows(content, sheet)
    header_at = _header_row(rows, "trimestre" if frequency == "Q" else "año", sheet)
    header = [c for c in rows[header_at] if c is not None]
    # "Número de giros1": the trailing digit is a footnote marker glued to the label.
    keys = [re.sub(r"_?\d+$", "", snake_case(h)) for h in header[1:]]
    columns = [TOTAL_COLUMNS.get(k, k) for k in keys]
    offset = rows[header_at].index(header[0])

    periods, values = [], []
    for row in rows[header_at + 1 :]:
        label = row[offset] if len(row) > offset else None
        try:
            period = parse_period(label)
        except (ValueError, TypeError):
            if periods:
                break  # footnotes
            continue
        periods.append(period)
        cells = row[offset + 1 : offset + 1 + len(columns)]
        values.append([to_number(c) for c in cells] + [float("nan")] * (len(columns) - len(cells)))
    if not periods:
        raise ParseError(f"La hoja {sheet!r} no tiene filas de datos.")
    index = pd.PeriodIndex(periods, freq=freq, name="period")
    frame = pd.DataFrame(values, index=index, columns=columns, dtype="float64")
    if "monto_usd" not in frame.columns:
        raise ParseError(
            f"A la hoja {sheet!r} le falta la columna 'Monto en USD'; se encontró {header}."
        )
    if not frame.index.is_unique or not frame.index.is_monotonic_increasing:
        raise ParseError(f"Los períodos de la hoja {sheet!r} están repetidos o desordenados.")
    frame.attrs["original_columns"] = dict(
        zip(columns, (str(h).strip() for h in header[1:]), strict=True)
    )
    return frame


def parse_by_country(content: bytes, sheet: str, frequency: str) -> pd.DataFrame:
    """``RRpais_t``-style sheet -> DataFrame[period] with one column per country plus ``total``."""
    _, parse_period, freq = FREQUENCIES[frequency]
    rows = _rows(content, sheet)
    header_at = _header_row(rows, "código", sheet)
    header = rows[header_at]
    offset = next(i for i, c in enumerate(header) if c is not None)
    period_cells = [c for c in header[offset + 2 :] if c is not None]
    try:
        periods = [parse_period(c) for c in period_cells]
    except ValueError as exc:
        raise ParseError(f"Etiqueta de período no reconocida en la hoja {sheet!r}: {exc}") from exc

    data: dict[str, list[float]] = {}
    names: dict[str, str] = {}
    codes: dict[str, str | None] = {}
    merged: list[str] = []
    for row in rows[header_at + 1 :]:
        code = row[offset] if len(row) > offset else None
        label = row[offset + 1] if len(row) > offset + 1 else None
        if code is None and label is None:
            if data:
                break
            continue
        if str(code).strip().lower() == "total":
            column, name, code = "total", "Total", None
        else:
            name = str(label).strip()
            column = snake_case(name)
            code = str(code).strip()
        cells = list(row[offset + 2 : offset + 2 + len(periods)])
        cells += [None] * (len(periods) - len(cells))
        values = [to_number(c) for c in cells]
        # BCE sometimes splits one country over two rows that differ only in capitalisation
        # ("San Vicente y Las Granadinas" / "... las Granadinas"); merge rows sharing a code.
        existing = next((c for c, k in codes.items() if k is not None and k == code), None)
        if existing is not None:
            merged.append(name)
            data[existing] = [_combine(a, b) for a, b in zip(data[existing], values, strict=True)]
            continue
        if column in data:
            raise ParseError(f"El país {name!r} aparece dos veces en la hoja {sheet!r}.")
        data[column] = values
        names[column] = name
        codes[column] = code
    if "total" not in data:
        raise ParseError(f"La hoja {sheet!r} no tiene fila 'Total'.")
    index = pd.PeriodIndex(periods, freq=freq, name="period")
    frame = pd.DataFrame(data, index=index, dtype="float64")
    if not frame.index.is_unique or not frame.index.is_monotonic_increasing:
        raise ParseError(f"Los períodos de la hoja {sheet!r} están repetidos o desordenados.")
    frame.attrs["original_columns"] = names
    frame.attrs["country_codes"] = codes
    if merged:
        frame.attrs["merged_duplicate_rows"] = merged
    return frame


def _combine(a: float, b: float) -> float:
    if a != a:  # NaN
        return b
    if b != b:
        return a
    return a + b


def remesas(
    flow: str = "received",
    frequency: str = "Q",
    by: str | None = None,
    count: bool = False,
    refresh: bool = False,
) -> pd.DataFrame:
    """Remesas de trabajadores recibidas por el Ecuador, o enviadas desde el Ecuador.

    Args:
        flow: ``"received"`` (por defecto; remesas recibidas, desde 1993) o ``"sent"``
            (remesas enviadas, desde 2005).
        frequency: ``"Q"`` trimestral (por defecto) o ``"A"`` anual.
        by: ``None`` para el total nacional, o ``"country"`` para una columna por país de
            origen (recibidas) o destino (enviadas) más ``total``. Los datos por país
            empiezan en 2007 (recibidas) y 2009Q4 (enviadas).
        count: Con ``by="country"``, devolver el número de giros en lugar del monto.
        refresh: Volver a descargar aunque haya una copia vigente en el caché.

    Returns:
        DataFrame con ``PeriodIndex`` trimestral o anual llamado ``period``. Los totales
        tienen ``monto_usd`` (USD) y ``numero_giros`` (número de giros, desde 2007/2010).
        ``df.attrs`` contiene ``source``, ``dataset``, ``flow``, ``units``, ``url``,
        ``retrieved_at``, ``original_columns`` y, por país, ``country_codes``.

    Raises:
        ValueError: combinación de argumentos inválida.
        SourceUnavailableError / OfflineError / ParseError.

    Example:
        >>> r = ec.bce.remesas()
        >>> r.loc["2024Q1":, "monto_usd"] / 1e6       # millones de USD
        >>> ec.bce.remesas(by="country", frequency="A")[["estados_unidos_de_america", "espana"]]
    """
    sheet = sheet_name(flow, frequency, by, count)  # validates arguments before downloading
    workbook = fetch(refresh=refresh)
    if by == "country":
        frame = parse_by_country(workbook.content, sheet, frequency)
        units = "número de giros" if count else "USD"
    else:
        frame = parse_totals(workbook.content, sheet, frequency)
        units = "USD (monto_usd), número de giros (numero_giros)"
    frame.attrs.update(
        {
            "source": SOURCE,
            "dataset": "remesas",
            "flow": flow,
            "frequency": frequency,
            "by": by,
            "sheet": sheet,
            "units": units,
            "url": workbook.url,
            "landing_url": LANDING_URL,
            "retrieved_at": workbook.retrieved_at.isoformat(),
        }
    )
    return frame
